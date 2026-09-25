"""Per-category monthly budgets: defaults, one-month overrides, and where they show up."""

import io

import app as vb_app
import budget

from .conftest import csrf_token, register_and_login


def _user(username):
    conn = vb_app.get_connection()
    uid = conn.execute("SELECT id FROM users WHERE username = ?", (username,)).fetchone()["id"]
    cats = {
        row["name"]: row["id"]
        for row in conn.execute("SELECT id, name FROM categories WHERE user_id = ?", (uid,))
    }
    account_id = conn.execute(
        "SELECT id FROM accounts WHERE user_id = ? ORDER BY id LIMIT 1", (uid,)
    ).fetchone()["id"]
    conn.close()
    return uid, cats, account_id


def _spend(uid, category_id, account_id, amount, day):
    conn = vb_app.get_connection()
    conn.execute(
        "INSERT INTO expenses (user_id, notes, amount, category_id, account_id, spent_at)"
        " VALUES (?, '', ?, ?, ?, ?)",
        (uid, -abs(amount), category_id, account_id, day),
    )
    conn.commit()
    conn.close()


def _refund(uid, category_id, account_id, amount, day):
    """Insert a refund: stored amount is positive (spent = -SUM(amount))."""
    conn = vb_app.get_connection()
    conn.execute(
        "INSERT INTO expenses (user_id, notes, amount, category_id, account_id, spent_at)"
        " VALUES (?, '', ?, ?, ?, ?)",
        (uid, abs(amount), category_id, account_id, day),
    )
    conn.commit()
    conn.close()


def _post(client, url, data):
    return client.post(url, data={**data, "_csrf_token": csrf_token(client)}, follow_redirects=True)


def _summary(uid, ym):
    conn = vb_app.get_connection()
    try:
        return budget.budget_for_month(conn, uid, ym)
    finally:
        conn.close()


def test_toggle_enables_and_disables(client):
    register_and_login(client, "budgettoggle")
    uid, _, _ = _user("budgettoggle")

    page = client.get("/?panel=budget").get_data(as_text=True)
    assert "Enable budgeting" in page

    _post(client, "/budget/toggle", {"enabled": "1"})
    conn = vb_app.get_connection()
    assert budget.is_enabled(conn, uid)
    conn.close()
    page = client.get("/?panel=budget").get_data(as_text=True)
    assert "Monthly budgets" in page

    _post(client, "/budget/toggle", {"enabled": "0"})
    conn = vb_app.get_connection()
    assert not budget.is_enabled(conn, uid)
    conn.close()


def test_default_and_override_drive_the_month(client):
    register_and_login(client, "budgetmath")
    uid, cats, account_id = _user("budgetmath")
    general, other = cats["General"], cats["Other"]
    _post(client, "/budget/toggle", {"enabled": "1"})
    _post(client, "/budget/defaults", {f"amount_{general}": "100", f"amount_{other}": ""})

    _spend(uid, general, account_id, 85, "2026-03-10")
    _spend(uid, general, account_id, 30, "2026-04-02")
    _spend(uid, other, account_id, 12, "2026-03-11")

    march = _summary(uid, "2026-03")
    row = march["by_category"][general]
    assert row["budget"] == 100 and row["spent"] == 85 and row["remaining"] == 15
    assert row["status"] == "warn"
    # Other has no budget: listed separately and kept out of the totals.
    assert other not in march["by_category"]
    assert [r["name"] for r in march["unbudgeted"]] == ["Other"]
    assert march["total_budget"] == 100 and march["total_spent"] == 85

    _post(client, "/budget/override", {"category_id": str(general), "month": "2026-03", "amount": "80"})
    march = _summary(uid, "2026-03")
    assert march["by_category"][general]["budget"] == 80
    assert march["by_category"][general]["status"] == "over"
    # The override is for March only.
    april = _summary(uid, "2026-04")
    assert april["by_category"][general]["budget"] == 100
    assert april["by_category"][general]["status"] == "ok"

    _post(client, "/budget/override", {"category_id": str(general), "month": "2026-03", "clear": "1"})
    assert _summary(uid, "2026-03")["by_category"][general]["budget"] == 100

    # Blank default removes the budget.
    _post(client, "/budget/defaults", {f"amount_{general}": ""})
    assert _summary(uid, "2026-04")["rows"] == []


def test_negative_amount_is_rejected_without_partial_save(client):
    register_and_login(client, "budgetbad")
    uid, cats, _ = _user("budgetbad")
    _post(client, "/budget/toggle", {"enabled": "1"})
    resp = _post(
        client,
        "/budget/defaults",
        {f"amount_{cats['General']}": "50", f"amount_{cats['Other']}": "-5"},
    )
    assert "zero or more" in resp.get_data(as_text=True)
    conn = vb_app.get_connection()
    assert budget.default_budgets(conn, uid) == {}
    conn.close()


def test_cannot_budget_someone_elses_category(client, app):
    register_and_login(client, "budgetowner")
    _, owner_cats, _ = _user("budgetowner")

    intruder = app.test_client()
    register_and_login(intruder, "budgetintruder")
    intruder_uid, _, _ = _user("budgetintruder")
    _post(intruder, "/budget/override", {"category_id": str(owner_cats["General"]), "month": "2026-03", "amount": "5"})
    _post(intruder, "/budget/defaults", {f"amount_{owner_cats['General']}": "5"})

    conn = vb_app.get_connection()
    assert not conn.execute(
        "SELECT 1 FROM category_budget_overrides WHERE category_id = ?", (owner_cats["General"],)
    ).fetchone()
    assert not conn.execute(
        "SELECT 1 FROM category_budgets WHERE category_id = ?", (owner_cats["General"],)
    ).fetchone()
    assert budget.default_budgets(conn, intruder_uid) == {}
    conn.close()


def test_deleting_a_category_drops_its_budgets(client):
    register_and_login(client, "budgetcascade")
    uid, _, _ = _user("budgetcascade")
    _post(client, "/categories/add", {"name": "Travel"})
    _, cats, _ = _user("budgetcascade")
    travel = cats["Travel"]
    _post(client, "/budget/defaults", {f"amount_{travel}": "200"})
    _post(client, "/budget/override", {"category_id": str(travel), "month": "2026-08", "amount": "900"})

    _post(client, f"/categories/{travel}/delete", {})

    conn = vb_app.get_connection()
    assert not conn.execute("SELECT 1 FROM categories WHERE id = ?", (travel,)).fetchone()
    assert not conn.execute("SELECT 1 FROM category_budgets WHERE category_id = ?", (travel,)).fetchone()
    assert not conn.execute(
        "SELECT 1 FROM category_budget_overrides WHERE category_id = ?", (travel,)
    ).fetchone()
    conn.close()


def test_home_and_consolidated_show_budget_when_enabled(client):
    register_and_login(client, "budgetviews")
    uid, cats, account_id = _user("budgetviews")
    general = cats["General"]
    _post(client, "/budget/defaults", {f"amount_{general}": "50"})
    _spend(uid, general, account_id, 70, "2026-02-03")

    # Amounts are kept while budgeting is off, but nothing is shown.
    page = client.get("/?panel=home&month=2026-02").get_data(as_text=True)
    assert "home-budget-card" not in page

    _post(client, "/budget/toggle", {"enabled": "1"})
    home = client.get("/?panel=home&month=2026-02").get_data(as_text=True)
    assert "home-budget-card" in home
    assert "Over budget: General" in home

    summary = client.get("/?panel=summary&month=2026-02").get_data(as_text=True)
    assert "<th>Budget</th>" in summary and "<th>Left</th>" in summary


def test_budgets_survive_export_import(client, app):
    register_and_login(client, "budgetexport")
    _, cats, _ = _user("budgetexport")
    _post(client, "/categories/add", {"name": "Groceries"})
    _, cats, _ = _user("budgetexport")
    _post(client, "/budget/toggle", {"enabled": "1"})
    _post(client, "/budget/defaults", {f"amount_{cats['Groceries']}": "300"})
    _post(client, "/budget/override", {"category_id": str(cats["Groceries"]), "month": "2026-12", "amount": "450"})
    workbook_bytes = client.get("/export/excel").data

    other = app.test_client()
    register_and_login(other, "budgetimport")
    resp = other.post(
        "/import/excel",
        data={
            "file": (io.BytesIO(workbook_bytes), "budget-export.xlsx"),
            "replace_movements": "1",
            "_csrf_token": csrf_token(other),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert b"Import failed" not in resp.data

    dst_uid, dst_cats, _ = _user("budgetimport")
    conn = vb_app.get_connection()
    assert budget.is_enabled(conn, dst_uid)
    assert budget.default_budgets(conn, dst_uid) == {dst_cats["Groceries"]: 300.0}
    conn.close()
    assert _summary(dst_uid, "2026-12")["by_category"][dst_cats["Groceries"]]["budget"] == 450


# ---------------------------------------------------------------------------
# spending_history
# ---------------------------------------------------------------------------


def test_spending_history_average_window_and_refunds(client):
    register_and_login(client, "histwindow")
    uid, cats, account_id = _user("histwindow")
    general = cats["General"]

    # Spending in the selected month itself must be ignored.
    _spend(uid, general, account_id, 999, "2026-05-15")
    # 3-month window before 2026-05: Feb, Mar, Apr.
    _spend(uid, general, account_id, 90, "2026-02-10")
    _spend(uid, general, account_id, 60, "2026-03-10")
    _spend(uid, general, account_id, 30, "2026-04-10")
    # 4 months back (Jan) must be ignored.
    _spend(uid, general, account_id, 1000, "2026-01-10")

    history = budget.spending_history(vb_conn(uid), uid, "2026-05")
    row = history[general]
    assert row["avg"] == 60.0  # (90+60+30)/3
    assert row["suggested"] == 60.0


def test_spending_history_refunds_reduce_average_and_clamp_at_zero(client):
    register_and_login(client, "histrefund")
    uid, cats, account_id = _user("histrefund")
    general = cats["General"]

    _spend(uid, general, account_id, 90, "2026-02-10")
    _refund(uid, general, account_id, 200, "2026-03-10")  # refund exceeds spending
    _spend(uid, general, account_id, 30, "2026-04-10")

    history = budget.spending_history(vb_conn(uid), uid, "2026-05")
    row = history[general]
    # Net = 90 - 200 + 30 = -80, clamped to 0.
    assert row["avg"] == 0.0
    assert row["suggested"] == 0.0


def test_spending_history_new_user_divides_by_months_actually_lived(client):
    register_and_login(client, "histnew")
    uid, cats, account_id = _user("histnew")
    general = cats["General"]

    # First-ever expense was last month (April) relative to selected month May.
    _spend(uid, general, account_id, 100, "2026-04-15")

    history = budget.spending_history(vb_conn(uid), uid, "2026-05")
    row = history[general]
    # Only April counts (Feb/Mar predate the user's first expense) -> divide by 1.
    assert row["avg"] == 100.0
    assert row["suggested"] == 100.0


def test_spending_history_suggested_rounds_up_to_multiple_of_five(client):
    register_and_login(client, "histround")
    uid, cats, account_id = _user("histround")
    general = cats["General"]

    # Anchor the user's first-ever expense before the 3-month window so the
    # whole window counts, rather than being divided by fewer months.
    anchor = cats["Other"]
    _spend(uid, anchor, account_id, 0.01, "2026-01-05")

    _spend(uid, general, account_id, 633, "2026-04-01")  # avg 211 -> suggest 215

    history = budget.spending_history(vb_conn(uid), uid, "2026-05")
    row = history[general]
    assert row["avg"] == 211.0
    assert row["suggested"] == 215.0

    conn = vb_app.get_connection()
    row = conn.execute(
        "INSERT INTO categories (user_id, name) VALUES (?, 'Round215')", (uid,)
    )
    other = row.lastrowid
    conn.commit()
    conn.close()
    _spend(uid, other, account_id, 645, "2026-04-01")  # avg 215 -> stays 215 (window fully counted via anchor)
    history = budget.spending_history(vb_conn(uid), uid, "2026-05")
    assert history[other]["avg"] == 215.0
    assert history[other]["suggested"] == 215.0

    # A category with no spending at all: avg 0 stays suggested 0.
    conn = vb_app.get_connection()
    conn.execute("INSERT INTO categories (user_id, name) VALUES (?, 'Untouched')", (uid,))
    conn.commit()
    conn.close()
    history = budget.spending_history(vb_conn(uid), uid, "2026-05")
    untouched_id = [row["id"] for row in vb_conn(uid).execute(
        "SELECT id FROM categories WHERE user_id = ? AND name = 'Untouched'", (uid,)
    )][0]
    assert history[untouched_id]["avg"] == 0.0
    assert history[untouched_id]["suggested"] == 0.0


def test_spending_history_active_flags(client, app):
    register_and_login(client, "histactive")
    uid, cats, account_id = _user("histactive")
    general, other = cats["General"], cats["Other"]

    # General: spent 7 months before May 2026 (October 2025) -> inactive by recency.
    _spend(uid, general, account_id, 50, "2025-10-05")
    # Other: no spending at all, but has a default budget -> active.
    conn = vb_app.get_connection()
    budget.set_default(conn, uid, other, 20)
    conn.commit()
    conn.close()

    history = budget.spending_history(vb_conn(uid), uid, "2026-05")
    assert history[general]["active"] is False
    assert history[other]["active"] is True

    # An override for the selected month also makes a category active.
    conn = vb_app.get_connection()
    budget.set_override(conn, uid, general, "2026-05", 15)
    conn.commit()
    conn.close()
    history = budget.spending_history(vb_conn(uid), uid, "2026-05")
    assert history[general]["active"] is True

    # Recent spending (within the last 6 months) also makes a category active.
    other_client = app.test_client()
    register_and_login(other_client, "histactive2")
    uid2, cats2, account_id2 = _user("histactive2")
    general2 = cats2["General"]
    _spend(uid2, general2, account_id2, 10, "2025-12-01")  # 5 months before May 2026
    history2 = budget.spending_history(vb_conn(uid2), uid2, "2026-05")
    assert history2[general2]["active"] is True


# ---------------------------------------------------------------------------
# other_amount
# ---------------------------------------------------------------------------


def test_other_amount_set_reflected_in_totals_and_over(client):
    register_and_login(client, "otherbasic")
    uid, cats, account_id = _user("otherbasic")
    general, other = cats["General"], cats["Other"]
    _post(client, "/budget/toggle", {"enabled": "1"})
    _post(client, "/budget/defaults", {f"amount_{general}": "100", "other_amount": "30"})

    _spend(uid, other, account_id, 40, "2026-05-05")  # Other has no own budget

    summary = _summary(uid, "2026-05")
    assert summary["other"]["budget"] == 30
    assert summary["other"]["spent"] == 40
    assert summary["other"]["status"] == "over"
    assert summary["total_budget"] == 130
    assert summary["total_spent"] == 40
    assert summary["other"] in summary["over"]


def test_other_amount_cleared_by_blank(client):
    register_and_login(client, "otherclear")
    uid, cats, _ = _user("otherclear")
    _post(client, "/budget/toggle", {"enabled": "1"})
    _post(client, "/budget/defaults", {"other_amount": "30"})
    conn = vb_app.get_connection()
    assert budget.get_other_amount(conn, uid) == 30.0
    conn.close()

    _post(client, "/budget/defaults", {"other_amount": ""})
    conn = vb_app.get_connection()
    assert budget.get_other_amount(conn, uid) is None
    conn.close()
    assert _summary(uid, "2026-05")["other"] is None


def test_other_amount_negative_rejected_without_partial_save(client):
    register_and_login(client, "otherneg")
    uid, cats, _ = _user("otherneg")
    general = cats["General"]
    _post(client, "/budget/toggle", {"enabled": "1"})
    _post(client, "/budget/defaults", {"other_amount": "20"})

    resp = _post(client, "/budget/defaults", {f"amount_{general}": "50", "other_amount": "-5"})
    assert "zero or more" in resp.get_data(as_text=True)

    conn = vb_app.get_connection()
    # Neither the new default nor the other_amount change was saved.
    assert budget.default_budgets(conn, uid) == {}
    assert budget.get_other_amount(conn, uid) == 20.0
    conn.close()


def test_other_amount_works_without_prior_budget_settings_row(client):
    """A user who never toggled budgeting has no budget_settings row yet."""
    register_and_login(client, "othernorow")
    uid, cats, _ = _user("othernorow")
    conn = vb_app.get_connection()
    assert conn.execute(
        "SELECT 1 FROM budget_settings WHERE user_id = ?", (uid,)
    ).fetchone() is None
    conn.close()

    _post(client, "/budget/defaults", {"other_amount": "25"})

    conn = vb_app.get_connection()
    assert budget.get_other_amount(conn, uid) == 25.0
    assert not budget.is_enabled(conn, uid)  # enabled flag untouched
    conn.close()


# ---------------------------------------------------------------------------
# copy-last-month
# ---------------------------------------------------------------------------


def test_copy_last_month_creates_and_overwrites_overrides(client):
    register_and_login(client, "copybasic")
    uid, cats, account_id = _user("copybasic")
    general, other = cats["General"], cats["Other"]
    _post(client, "/budget/toggle", {"enabled": "1"})

    _spend(uid, general, account_id, 123.456, "2026-04-10")
    _refund(uid, other, account_id, 50, "2026-04-11")  # refund only, net <= 0

    # Existing override for May should be overwritten.
    conn = vb_app.get_connection()
    budget.set_override(conn, uid, general, "2026-05", 999)
    conn.commit()
    conn.close()

    resp = _post(client, "/budget/copy-last-month", {"month": "2026-05"})
    assert "Copied last month" in resp.get_data(as_text=True)

    conn = vb_app.get_connection()
    row = conn.execute(
        "SELECT amount FROM category_budget_overrides WHERE user_id = ? AND category_id = ? AND ym = ?",
        (uid, general, "2026-05"),
    ).fetchone()
    assert round(row["amount"], 2) == 123.46
    # Other only had a refund, so it should not be touched.
    assert conn.execute(
        "SELECT 1 FROM category_budget_overrides WHERE user_id = ? AND category_id = ? AND ym = ?",
        (uid, other, "2026-05"),
    ).fetchone() is None
    conn.close()


def test_copy_last_month_january_uses_december_previous_year(client):
    register_and_login(client, "copyjan")
    uid, cats, account_id = _user("copyjan")
    general = cats["General"]
    _post(client, "/budget/toggle", {"enabled": "1"})
    _spend(uid, general, account_id, 77, "2025-12-20")

    resp = _post(client, "/budget/copy-last-month", {"month": "2026-01"})
    assert "Copied last month" in resp.get_data(as_text=True)

    conn = vb_app.get_connection()
    row = conn.execute(
        "SELECT amount FROM category_budget_overrides WHERE user_id = ? AND category_id = ? AND ym = ?",
        (uid, general, "2026-01"),
    ).fetchone()
    assert row["amount"] == 77.0
    conn.close()


def test_copy_last_month_empty_previous_month_copies_nothing(client):
    register_and_login(client, "copyempty")
    uid, cats, _ = _user("copyempty")
    _post(client, "/budget/toggle", {"enabled": "1"})

    resp = _post(client, "/budget/copy-last-month", {"month": "2026-05"})
    assert "No spending last month to copy" in resp.get_data(as_text=True)

    conn = vb_app.get_connection()
    assert conn.execute(
        "SELECT 1 FROM category_budget_overrides WHERE user_id = ?", (uid,)
    ).fetchone() is None
    conn.close()


# ---------------------------------------------------------------------------
# Dashboard rendering
# ---------------------------------------------------------------------------


def test_dashboard_budget_panel_renders_form_controls_and_order(client):
    register_and_login(client, "dashbudget")
    uid, cats, account_id = _user("dashbudget")
    general, other_cat = cats["General"], cats["Other"]
    _post(client, "/budget/toggle", {"enabled": "1"})

    # General is active (recent spending); Other stays inactive (old spending only).
    _spend(uid, general, account_id, 50, "2026-04-10")
    _spend(uid, other_cat, account_id, 20, "2025-01-10")

    resp = client.get("/?panel=budget&month=2026-05")
    assert resp.status_code == 200
    page = resp.get_data(as_text=True)

    assert 'id="budget-search-input"' in page
    assert 'id="budget-fill-history-btn"' in page
    assert 'id="budget-distribute-btn"' in page
    assert 'name="other_amount"' in page
    # All categories now live together in a single collapsible section.
    assert 'id="budget-manual"' in page
    assert 'id="budget-inactive-details"' not in page

    general_pos = page.index(f'name="amount_{general}"')
    other_pos = page.index(f'name="amount_{other_cat}"')
    assert general_pos < other_pos  # active category (General) appears before inactive (Other)


# ---------------------------------------------------------------------------
# Export / import round trip for other_amount
# ---------------------------------------------------------------------------


def test_other_amount_survives_export_import_round_trip(client, app):
    register_and_login(client, "otherexport")
    uid, cats, _ = _user("otherexport")
    _post(client, "/budget/toggle", {"enabled": "1"})
    _post(client, "/budget/defaults", {"other_amount": "42"})
    workbook_bytes = client.get("/export/excel").data

    other = app.test_client()
    register_and_login(other, "otherimport")
    resp = other.post(
        "/import/excel",
        data={
            "file": (io.BytesIO(workbook_bytes), "budget-export.xlsx"),
            "replace_movements": "1",
            "_csrf_token": csrf_token(other),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert b"Import failed" not in resp.data

    dst_uid, _, _ = _user("otherimport")
    conn = vb_app.get_connection()
    assert budget.get_other_amount(conn, dst_uid) == 42.0
    conn.close()


def test_old_workbook_without_other_amount_meta_key_does_not_clear_it(client, app):
    """An older-format workbook has no `budget_other_amount` meta row; importing it
    over an account that already has an other_amount set should leave it alone."""
    register_and_login(client, "otherkeep")
    uid, cats, _ = _user("otherkeep")
    _post(client, "/budget/toggle", {"enabled": "1"})
    _post(client, "/budget/defaults", {"other_amount": "60"})

    workbook_bytes = client.get("/export/excel").data

    from openpyxl import load_workbook

    buf = io.BytesIO(workbook_bytes)
    wb = load_workbook(buf)
    ws_meta = wb["_meta"]
    # Remove the budget_other_amount row to simulate an older export.
    for row_idx in range(ws_meta.max_row, 1, -1):
        if str(ws_meta.cell(row=row_idx, column=1).value or "").strip() == "budget_other_amount":
            ws_meta.delete_rows(row_idx)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    resp = client.post(
        "/import/excel",
        data={
            "file": (out, "old-format.xlsx"),
            "replace_movements": "1",
            "_csrf_token": csrf_token(client),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert b"Import failed" not in resp.data

    conn = vb_app.get_connection()
    assert budget.get_other_amount(conn, uid) == 60.0
    conn.close()


def vb_conn(_uid=None):
    """Fresh connection for direct budget.* calls in these tests."""
    return vb_app.get_connection()


# ---------------------------------------------------------------------------
# Regression tests for code-review fixes
# ---------------------------------------------------------------------------


def test_spending_history_active_when_only_spending_is_in_selected_month(client):
    register_and_login(client, "histcurmonth")
    uid, cats, account_id = _user("histcurmonth")
    general = cats["General"]

    # Nothing in the 3-month window before May, but spending in May itself.
    _spend(uid, general, account_id, 40, "2026-05-05")

    history = budget.spending_history(vb_conn(), uid, "2026-05")
    row = history[general]
    assert row["active"] is True
    assert row["avg"] == 0.0
    assert row["suggested"] == 0.0


def test_budget_for_month_float_rounding_avoids_false_over(client):
    register_and_login(client, "floatround")
    uid, cats, account_id = _user("floatround")
    general = cats["General"]
    _post(client, "/budget/toggle", {"enabled": "1"})
    _post(client, "/budget/defaults", {f"amount_{general}": "0.30"})

    _spend(uid, general, account_id, 0.10, "2026-05-01")
    _spend(uid, general, account_id, 0.20, "2026-05-02")

    summary = _summary(uid, "2026-05")
    row = summary["by_category"][general]
    assert row["spent"] == 0.30
    assert row["status"] == "warn"
    assert row["remaining"] == 0.0
    assert summary["total_budget"] == 0.30
    assert summary["total_spent"] == 0.30
    assert summary["total_remaining"] == 0.0


def test_copy_last_month_with_other_amount_skips_unbudgeted_categories(client):
    register_and_login(client, "copyother")
    uid, cats, account_id = _user("copyother")
    general, other_cat = cats["General"], cats["Other"]
    _post(client, "/budget/toggle", {"enabled": "1"})
    # General has its own default budget; Other has none, so it is covered by
    # "everything else" and must not get its own override.
    _post(client, "/budget/defaults", {f"amount_{general}": "100", "other_amount": "50"})

    _spend(uid, general, account_id, 30, "2026-04-10")
    _spend(uid, other_cat, account_id, 15, "2026-04-11")

    resp = _post(client, "/budget/copy-last-month", {"month": "2026-05"})
    assert "Copied last month" in resp.get_data(as_text=True)

    conn = vb_app.get_connection()
    assert conn.execute(
        "SELECT amount FROM category_budget_overrides WHERE user_id = ? AND category_id = ? AND ym = ?",
        (uid, general, "2026-05"),
    ).fetchone()["amount"] == 30.0
    assert conn.execute(
        "SELECT 1 FROM category_budget_overrides WHERE user_id = ? AND category_id = ? AND ym = ?",
        (uid, other_cat, "2026-05"),
    ).fetchone() is None
    conn.close()


def test_copy_last_month_without_other_amount_copies_all_spending_categories(client):
    register_and_login(client, "copynoother")
    uid, cats, account_id = _user("copynoother")
    general, other_cat = cats["General"], cats["Other"]
    _post(client, "/budget/toggle", {"enabled": "1"})
    _post(client, "/budget/defaults", {f"amount_{general}": "100"})  # no other_amount

    _spend(uid, general, account_id, 30, "2026-04-10")
    _spend(uid, other_cat, account_id, 15, "2026-04-11")

    resp = _post(client, "/budget/copy-last-month", {"month": "2026-05"})
    assert "Copied last month" in resp.get_data(as_text=True)

    conn = vb_app.get_connection()
    assert conn.execute(
        "SELECT amount FROM category_budget_overrides WHERE user_id = ? AND category_id = ? AND ym = ?",
        (uid, general, "2026-05"),
    ).fetchone()["amount"] == 30.0
    assert conn.execute(
        "SELECT amount FROM category_budget_overrides WHERE user_id = ? AND category_id = ? AND ym = ?",
        (uid, other_cat, "2026-05"),
    ).fetchone()["amount"] == 15.0
    conn.close()


def test_parse_amount_rejects_non_finite_and_over_ceiling():
    for junk in ("inf", "Infinity", "-inf", "1e999", "nan"):
        try:
            budget.parse_amount(junk)
            raise AssertionError(f"expected ValueError for {junk!r}")
        except ValueError:
            pass
    try:
        budget.parse_amount(str(budget.MAX_AMOUNT + 1))
        raise AssertionError("expected ValueError above MAX_AMOUNT")
    except ValueError:
        pass
    # Exactly at the ceiling is fine.
    assert budget.parse_amount(str(budget.MAX_AMOUNT)) == budget.MAX_AMOUNT


def test_defaults_route_rejects_junk_amounts_without_partial_write(client):
    register_and_login(client, "junkdefaults")
    uid, cats, _ = _user("junkdefaults")
    general, other_cat = cats["General"], cats["Other"]
    _post(client, "/budget/toggle", {"enabled": "1"})

    for junk in ("inf", "Infinity", "1e999", "nan", str(budget.MAX_AMOUNT + 1)):
        resp = _post(
            client,
            "/budget/defaults",
            {f"amount_{general}": "50", f"amount_{other_cat}": junk},
        )
        assert resp.status_code == 200
        conn = vb_app.get_connection()
        assert budget.default_budgets(conn, uid) == {}, f"partial write for {junk!r}"
        conn.close()


def test_override_route_rejects_junk_amounts(client):
    register_and_login(client, "junkoverride")
    uid, cats, _ = _user("junkoverride")
    general = cats["General"]
    _post(client, "/budget/toggle", {"enabled": "1"})

    for junk in ("inf", "Infinity", "1e999", "nan", str(budget.MAX_AMOUNT + 1)):
        resp = _post(
            client,
            "/budget/override",
            {"category_id": str(general), "month": "2026-05", "amount": junk},
        )
        assert resp.status_code == 200
        conn = vb_app.get_connection()
        assert conn.execute(
            "SELECT 1 FROM category_budget_overrides WHERE user_id = ? AND category_id = ?",
            (uid, general),
        ).fetchone() is None
        conn.close()


def test_import_without_budgets_sheet_keeps_existing_budgets(client, app):
    register_and_login(client, "importnosheet")
    uid, cats, _ = _user("importnosheet")
    _post(client, "/budget/toggle", {"enabled": "1"})
    _post(client, "/budget/defaults", {f"amount_{cats['General']}": "77"})
    _post(client, "/budget/override", {"category_id": str(cats["General"]), "month": "2026-05", "amount": "88"})

    workbook_bytes = client.get("/export/excel").data

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(workbook_bytes))
    del wb["Budgets"]
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    resp = client.post(
        "/import/excel",
        data={
            "file": (out, "no-budgets-sheet.xlsx"),
            "replace_movements": "1",
            "_csrf_token": csrf_token(client),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert b"Import failed" not in resp.data

    conn = vb_app.get_connection()
    assert budget.default_budgets(conn, uid) == {cats["General"]: 77.0}
    assert conn.execute(
        "SELECT amount FROM category_budget_overrides WHERE user_id = ? AND category_id = ? AND ym = ?",
        (uid, cats["General"], "2026-05"),
    ).fetchone()["amount"] == 88.0
    conn.close()


def test_import_with_budgets_sheet_replaces_existing_budgets(client, app):
    register_and_login(client, "importwithsheet")
    uid, cats, _ = _user("importwithsheet")
    _post(client, "/budget/toggle", {"enabled": "1"})
    _post(client, "/budget/defaults", {f"amount_{cats['General']}": "77"})

    workbook_bytes = client.get("/export/excel").data

    # Change the default budget after exporting, then re-import the export —
    # replace mode with a Budgets sheet present must restore the exported value.
    _post(client, "/budget/defaults", {f"amount_{cats['General']}": "999"})

    resp = client.post(
        "/import/excel",
        data={
            "file": (io.BytesIO(workbook_bytes), "with-budgets-sheet.xlsx"),
            "replace_movements": "1",
            "_csrf_token": csrf_token(client),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert b"Import failed" not in resp.data

    conn = vb_app.get_connection()
    assert budget.default_budgets(conn, uid) == {cats["General"]: 77.0}
    conn.close()


def test_spending_history_today_param_limits_window_to_past_months(client):
    from datetime import date

    register_and_login(client, "historytoday")
    uid, cats, account_id = _user("historytoday")
    general = cats["General"]

    # ym = 2026-07, months=3 -> window Apr, May, Jun. today = 2026-05-15, so
    # only Apr and May are "real" months; Jun is in the future and must not
    # be counted in the denominator (even if it somehow has data).
    _spend(uid, general, account_id, 100, "2026-04-10")
    _spend(uid, general, account_id, 50, "2026-05-10")

    history = budget.spending_history(
        vb_conn(), uid, "2026-07", today=date(2026, 5, 15)
    )
    row = history[general]
    assert row["avg"] == 75.0  # (100+50)/2, not /3


def test_total_pct_is_100_when_all_budgets_zero_and_spending_exists(client):
    register_and_login(client, "zerobudget")
    uid, cats, account_id = _user("zerobudget")
    general = cats["General"]
    _post(client, "/budget/toggle", {"enabled": "1"})
    _post(client, "/budget/defaults", {f"amount_{general}": "0"})

    _spend(uid, general, account_id, 50, "2026-05-05")

    summary = _summary(uid, "2026-05")
    assert summary["total_budget"] == 0.0
    assert summary["total_spent"] == 50.0
    assert summary["total_pct"] == 100.0


def test_override_route_rejects_invalid_month(client):
    register_and_login(client, "invalidmonthoverride")
    uid, cats, _ = _user("invalidmonthoverride")
    general = cats["General"]
    _post(client, "/budget/toggle", {"enabled": "1"})

    for bad_month in ("2026-13", "abc", ""):
        resp = _post(
            client,
            "/budget/override",
            {"category_id": str(general), "month": bad_month, "amount": "10"},
        )
        assert "Invalid month." in resp.get_data(as_text=True)

    conn = vb_app.get_connection()
    assert conn.execute(
        "SELECT 1 FROM category_budget_overrides WHERE user_id = ?", (uid,)
    ).fetchone() is None
    conn.close()


def test_copy_last_month_rejects_invalid_month(client):
    register_and_login(client, "invalidmonthcopy")
    uid, cats, account_id = _user("invalidmonthcopy")
    general = cats["General"]
    _post(client, "/budget/toggle", {"enabled": "1"})
    _spend(uid, general, account_id, 20, "2026-04-10")

    for bad_month in ("2026-13", "abc", ""):
        resp = _post(client, "/budget/copy-last-month", {"month": bad_month})
        assert "Invalid month." in resp.get_data(as_text=True)

    conn = vb_app.get_connection()
    assert conn.execute(
        "SELECT 1 FROM category_budget_overrides WHERE user_id = ?", (uid,)
    ).fetchone() is None
    conn.close()


def test_home_card_renders_with_only_other_amount_set(client):
    register_and_login(client, "homeotheronly")
    uid, cats, account_id = _user("homeotheronly")
    _post(client, "/budget/toggle", {"enabled": "1"})
    _post(client, "/budget/defaults", {"other_amount": "40"})

    page = client.get("/?panel=home&month=2026-05").get_data(as_text=True)
    assert "home-budget-card" in page


def _flag_rows_for_category(category_id):
    conn = vb_app.get_connection()
    rows = conn.execute(
        "SELECT user_id, fixed FROM category_budget_flags WHERE category_id = ?",
        (category_id,),
    ).fetchall()
    conn.close()
    return rows


# ---------------------------------------------------------------------------
# "Fixed" per-category flag
# ---------------------------------------------------------------------------


def test_fixed_flag_persists_clears_and_is_untouched_when_amount_field_absent(client):
    register_and_login(client, "fixedbasic")
    uid, cats, _ = _user("fixedbasic")
    general, other_cat = cats["General"], cats["Other"]
    _post(client, "/budget/toggle", {"enabled": "1"})

    # Checking the box persists the flag even with a blank amount.
    _post(client, "/budget/defaults", {f"amount_{general}": "", f"fixed_{general}": "1"})
    conn = vb_app.get_connection()
    assert budget.fixed_categories(conn, uid) == {general}
    assert budget.default_budgets(conn, uid) == {}
    conn.close()

    # Submitting the same row's amount without the checkbox clears the flag.
    _post(client, "/budget/defaults", {f"amount_{general}": ""})
    conn = vb_app.get_connection()
    assert budget.fixed_categories(conn, uid) == set()
    conn.close()

    # Re-set it, then post a form that omits amount_<general> entirely
    # (only touches amount_<other_cat>) -- general's flag must be untouched.
    _post(client, "/budget/defaults", {f"amount_{general}": "", f"fixed_{general}": "1"})
    _post(client, "/budget/defaults", {f"amount_{other_cat}": "20"})
    conn = vb_app.get_connection()
    assert budget.fixed_categories(conn, uid) == {general}
    conn.close()


def test_fixed_flag_all_or_nothing_on_invalid_amount(client):
    register_and_login(client, "fixedallnone")
    uid, cats, _ = _user("fixedallnone")
    general, other_cat = cats["General"], cats["Other"]
    _post(client, "/budget/toggle", {"enabled": "1"})

    resp = _post(
        client,
        "/budget/defaults",
        {f"amount_{general}": "50", f"fixed_{general}": "1", f"amount_{other_cat}": "not-a-number"},
    )
    assert "zero or more" in resp.get_data(as_text=True)

    conn = vb_app.get_connection()
    assert budget.fixed_categories(conn, uid) == set()
    assert budget.default_budgets(conn, uid) == {}
    conn.close()


def test_fixed_flag_ignores_other_users_category(client, app):
    register_and_login(client, "fixedowner")
    _, owner_cats, _ = _user("fixedowner")
    owner_general = owner_cats["General"]

    intruder = app.test_client()
    register_and_login(intruder, "fixedintruder")
    _post(intruder, "/budget/toggle", {"enabled": "1"})
    _post(
        intruder,
        "/budget/defaults",
        {f"amount_{owner_general}": "10", f"fixed_{owner_general}": "1"},
    )

    assert _flag_rows_for_category(owner_general) == []


def test_deleting_category_cascades_fixed_flag(client):
    register_and_login(client, "fixedcascade")
    uid, _, _ = _user("fixedcascade")
    _post(client, "/categories/add", {"name": "Travel"})
    _, cats, _ = _user("fixedcascade")
    travel = cats["Travel"]
    _post(client, "/budget/toggle", {"enabled": "1"})
    _post(client, "/budget/defaults", {f"amount_{travel}": "", f"fixed_{travel}": "1"})

    conn = vb_app.get_connection()
    assert budget.fixed_categories(conn, uid) == {travel}
    conn.close()

    _post(client, f"/categories/{travel}/delete", {})

    conn = vb_app.get_connection()
    assert not conn.execute(
        "SELECT 1 FROM category_budget_flags WHERE category_id = ?", (travel,)
    ).fetchone()
    conn.close()


def test_copy_last_month_skips_fixed_categories(client):
    register_and_login(client, "fixedcopy")
    uid, cats, account_id = _user("fixedcopy")
    general, other_cat = cats["General"], cats["Other"]
    _post(client, "/budget/toggle", {"enabled": "1"})
    _post(
        client,
        "/budget/defaults",
        {f"amount_{general}": "100", f"fixed_{general}": "1", f"amount_{other_cat}": "50"},
    )

    _spend(uid, general, account_id, 30, "2026-04-10")
    _spend(uid, other_cat, account_id, 15, "2026-04-11")

    resp = _post(client, "/budget/copy-last-month", {"month": "2026-05"})
    assert "Copied last month" in resp.get_data(as_text=True)

    conn = vb_app.get_connection()
    # Fixed category: no override created, default budget untouched.
    assert conn.execute(
        "SELECT 1 FROM category_budget_overrides WHERE user_id = ? AND category_id = ? AND ym = ?",
        (uid, general, "2026-05"),
    ).fetchone() is None
    assert budget.default_budgets(conn, uid)[general] == 100.0
    # Non-fixed category gets the override.
    assert conn.execute(
        "SELECT amount FROM category_budget_overrides WHERE user_id = ? AND category_id = ? AND ym = ?",
        (uid, other_cat, "2026-05"),
    ).fetchone()["amount"] == 15.0
    conn.close()


def test_copy_last_month_skips_fixed_categories_with_other_amount_set(client):
    register_and_login(client, "fixedcopyother")
    uid, cats, account_id = _user("fixedcopyother")
    general, other_cat = cats["General"], cats["Other"]
    _post(client, "/budget/toggle", {"enabled": "1"})
    # General is fixed (own default); Other has no default -> covered by "everything else".
    _post(
        client,
        "/budget/defaults",
        {f"amount_{general}": "100", f"fixed_{general}": "1", "other_amount": "50"},
    )

    _spend(uid, general, account_id, 30, "2026-04-10")
    _spend(uid, other_cat, account_id, 15, "2026-04-11")

    resp = _post(client, "/budget/copy-last-month", {"month": "2026-05"})
    assert (
        "Nothing to copy — every category with spending last month is fixed "
        "or covered by Everything else." in resp.get_data(as_text=True)
    )

    conn = vb_app.get_connection()
    # General is fixed -> skipped. Other has no own budget -> skipped by the
    # other_amount rule. Nothing should have been copied.
    assert conn.execute(
        "SELECT 1 FROM category_budget_overrides WHERE user_id = ?", (uid,)
    ).fetchone() is None
    conn.close()


def test_budget_form_rows_expose_fixed_and_checkbox_is_checked(client):
    register_and_login(client, "fixedform")
    uid, cats, account_id = _user("fixedform")
    general, other_cat = cats["General"], cats["Other"]
    _post(client, "/budget/toggle", {"enabled": "1"})
    _spend(uid, general, account_id, 10, "2026-04-10")  # make it active/rendered
    _post(client, "/budget/defaults", {f"amount_{general}": "50", f"fixed_{general}": "1"})

    page = client.get("/?panel=budget&month=2026-05").get_data(as_text=True)
    marker = f'name="fixed_{general}"'
    assert marker in page
    pos = page.index(marker)
    tag_end = page.index(">", pos)
    assert "checked" in page[pos:tag_end]

    marker_other = f'name="fixed_{other_cat}"'
    assert marker_other in page
    pos_other = page.index(marker_other)
    tag_end_other = page.index(">", pos_other)
    assert "checked" not in page[pos_other:tag_end_other]


def test_excel_round_trip_preserves_fixed_flag_with_no_default_amount(client, app):
    register_and_login(client, "fixedexport")
    uid, cats, _ = _user("fixedexport")
    general = cats["General"]
    _post(client, "/budget/toggle", {"enabled": "1"})
    _post(client, "/budget/defaults", {f"amount_{general}": "", f"fixed_{general}": "1"})

    workbook_bytes = client.get("/export/excel").data

    other = app.test_client()
    register_and_login(other, "fixedimport")
    resp = other.post(
        "/import/excel",
        data={
            "file": (io.BytesIO(workbook_bytes), "fixed-export.xlsx"),
            "replace_movements": "1",
            "_csrf_token": csrf_token(other),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert b"Import failed" not in resp.data

    dst_uid, dst_cats, _ = _user("fixedimport")
    conn = vb_app.get_connection()
    assert budget.fixed_categories(conn, dst_uid) == {dst_cats["General"]}
    assert budget.default_budgets(conn, dst_uid) == {}
    conn.close()


def test_import_replace_mode_with_budgets_sheet_replaces_fixed_flag(client):
    register_and_login(client, "fixedreplace")
    uid, cats, _ = _user("fixedreplace")
    general = cats["General"]
    _post(client, "/budget/toggle", {"enabled": "1"})
    _post(client, "/budget/defaults", {f"amount_{general}": "50", f"fixed_{general}": "1"})

    workbook_bytes = client.get("/export/excel").data

    # Clear the flag locally.
    _post(client, "/budget/defaults", {f"amount_{general}": "50"})
    conn = vb_app.get_connection()
    assert budget.fixed_categories(conn, uid) == set()
    conn.close()

    # Re-importing the earlier export in replace mode restores the flag.
    resp = client.post(
        "/import/excel",
        data={
            "file": (io.BytesIO(workbook_bytes), "fixed-replace.xlsx"),
            "replace_movements": "1",
            "_csrf_token": csrf_token(client),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert b"Import failed" not in resp.data

    conn = vb_app.get_connection()
    assert budget.fixed_categories(conn, uid) == {general}
    conn.close()


def test_import_merge_mode_budgets_sheet_without_fixed_column_leaves_existing_flags(client):
    register_and_login(client, "fixedmerge")
    uid, cats, _ = _user("fixedmerge")
    general = cats["General"]
    _post(client, "/budget/toggle", {"enabled": "1"})
    _post(client, "/budget/defaults", {f"amount_{general}": "50", f"fixed_{general}": "1"})

    workbook_bytes = client.get("/export/excel").data

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(workbook_bytes))
    ws_bud = wb["Budgets"]
    # Blank out the "fixed" header so the importer treats the sheet as if it
    # predates the fixed column, while changing the amount to prove the
    # merge still applies the amount update.
    for col in range(1, ws_bud.max_column + 1):
        if str(ws_bud.cell(row=1, column=col).value or "").strip().lower() == "fixed":
            ws_bud.cell(row=1, column=col).value = None
    for row_idx in range(2, ws_bud.max_row + 1):
        if str(ws_bud.cell(row=row_idx, column=1).value or "").strip() == "General":
            ws_bud.cell(row=row_idx, column=3).value = 999  # amount column
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    resp = client.post(
        "/import/excel",
        data={
            "file": (out, "fixed-merge-no-column.xlsx"),
            "_csrf_token": csrf_token(client),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert b"Import failed" not in resp.data

    conn = vb_app.get_connection()
    assert budget.fixed_categories(conn, uid) == {general}
    assert budget.default_budgets(conn, uid) == {general: 999.0}
    conn.close()


def test_import_blank_amount_default_row_with_fixed_zero_is_an_error(client):
    register_and_login(client, "fixedbadrow")
    _user("fixedbadrow")

    from excel_io import _build_migration_template_workbook

    wb = _build_migration_template_workbook()
    ws_bud = wb["Budgets"]
    ws_bud.append(["General", "", "", 0])
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    resp = client.post(
        "/import/excel",
        data={
            "file": (out, "bad-budget-row.xlsx"),
            "_csrf_token": csrf_token(client),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert b"Import failed" in resp.data
    assert b"missing amount" in resp.data


def test_budget_panel_copy_last_month_confirm_and_hint_text(client):
    register_and_login(client, "panelconfirm")
    uid, cats, _ = _user("panelconfirm")
    _post(client, "/budget/toggle", {"enabled": "1"})

    page = client.get("/?panel=budget&month=2026-05").get_data(as_text=True)
    assert 'data-confirm=' in page
    assert "based on what you actually spent last month" in page
    assert "for budgeted categories" not in page

    _post(client, "/budget/defaults", {"other_amount": "40"})
    page = client.get("/?panel=budget&month=2026-05").get_data(as_text=True)
    assert "for budgeted categories" in page
    assert "based on what you actually spent last month" not in page
